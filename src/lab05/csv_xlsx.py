import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Конвертирует CSV в XLSX.
    Использовать openpyxl ИЛИ xlsxwriter.
    Первая строка CSV — заголовок.
    Лист называется "Sheet1".
    Колонки — автоширина по длине текста (не менее 8 символов).
    """

    csv_p = Path(csv_path)
    if not csv_p.is_file():
        raise FileNotFoundError("File not found")
    if csv_p.suffix.lower() != ".csv":
        raise ValueError("File must have .csv extension")

    with csv_p.open("r", newline="", encoding="utf-8") as file:
        reader = csv.reader(file)
        rows = list(reader)

    if not rows:
        raise ValueError("CSV file is empty")
    headers = rows[0]
    if not headers or any(not name.strip() for name in headers):
        raise ValueError("CSV headers must not be empty")
    if len(headers) != len(set(headers)):
        raise ValueError("CSV headers must be unique")
    if len(rows) < 2:
        raise ValueError("CSV must contain at least one data row")
    for row in rows:
        if len(row) != len(headers):
            raise ValueError("CSV file contains rows with missing or extra fields")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)

    for column in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        width = max(8, max_length + 2)

        letter = get_column_letter(column[0].column)
        ws.column_dimensions[letter].width = width

    xlsx_p = Path(xlsx_path)
    if xlsx_p.suffix.lower() != ".xlsx":
            raise ValueError("Output file must have .xlsx extension")
    
    xlsx_p.parent.mkdir(parents=True, exist_ok=True)
 

    wb.save(xlsx_p)
    
if __name__ == '__main__':
    csv_to_xlsx(r"data/samples/people.csv", r"data/out/people.xlsx")